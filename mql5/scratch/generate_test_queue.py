import os
import glob
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

xlsx_path = r"d:\FOR RENDO (ONLY )\Personal Project\Dashboard\mql5\Momentum_Backtest_Combined_V1.xlsx"
presets_dir = r"d:\FOR RENDO (ONLY )\Personal Project\Dashboard\mql5\Presets"

def parse_set_file(file_path):
    name = os.path.splitext(os.path.basename(file_path))[0]
    params = {}
    
    content = ""
    # Try reading as UTF-16 first (often used by MT5)
    try:
        with open(file_path, 'r', encoding='utf-16') as f:
            content = f.read()
    except Exception:
        # Try UTF-8 with ignore
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except Exception:
            # Fallback to Latin-1 which never fails on any byte sequence
            with open(file_path, 'r', encoding='latin-1') as f:
                content = f.read()

    for line in content.split('\n'):
        if '=' in line:
            key, val = line.split('=', 1)
            val_clean = val.split('||')[0].strip()
            params[key.strip()] = val_clean

    strategy_type = params.get('InpStrategyType', '4')
    strategy_map = {
        '0': 'ICT Sniper',
        '1': 'ICT Scalper',
        '2': 'Mean Reversion',
        '3': 'Breakout',
        '4': 'Momentum Scalper'
    }
    strategy = strategy_map.get(strategy_type, 'Unknown')

    pos_mode = params.get('InpPositionMode', '0')
    pos_map = {
        '0': 'Single',
        '1': 'Partial',
        '2': 'Grid'
    }
    position_mode = pos_map.get(pos_mode, 'Single')

    # Convert numeric helper values
    trailing = 'Yes' if params.get('InpTrailingStop', '0') in ['1', 'true', 'True'] else 'No'
    use_be = 'Yes' if params.get('InpUseBE', '0') in ['1', 'true', 'True'] else 'No'

    return {
        'PresetName': name,
        'Strategy': strategy,
        'Timeframe': params.get('InpSignalTimeframe', 'M1'),
        'Risk_Pct': params.get('InpRiskPercent', '1.0'),
        'MinRR': params.get('InpMinRR', '2.0'),
        'MaxSpread': params.get('InpMaxSpreadPoints', '200'),
        'Session': params.get('InpSession', 'both'),
        'PositionMode': position_mode,
        'MomLookback': params.get('InpMomLookbackSec', '-'),
        'MomThreshold': params.get('InpMomThresholdPips', '-'),
        'Cooldown_Sec': params.get('InpPyramidCooldownSec', '-'),
        'DailyLoss_Pct': params.get('InpMaxDailyLossPct', '5.0'),
        'MaxDD_Pct': params.get('InpMaxDrawdownPct', '15.0'),
        'TrailingStop': trailing,
        'UseBE': use_be
    }

# Read presets
presets = []
for file_path in glob.glob(os.path.join(presets_dir, "*.set")):
    try:
        presets.append(parse_set_file(file_path))
    except Exception as e:
        print(f"Skipping file {file_path} due to error: {e}")

# Load Workbook
wb = load_workbook(xlsx_path)

# Delete Test_Queue if exists
if "Test_Queue" in wb.sheetnames:
    del wb["Test_Queue"]

# Create at index 0
qs = wb.create_sheet("Test_Queue", 0)

headers = [
    "No", "Preset Name", "Strategy", "Timeframe", "Position Mode", "Risk %", "Min RR", 
    "Spread Limit", "Session", "Mom Lookback", "Mom Threshold", "Cooldown(s)", 
    "Daily Loss %", "Max DD %", "Trail Stop", "Break Even", "Status", 
    "Profit USD", "Profit Factor", "Equity DD %", "Trades", "Win Rate %", "Notes"
]

# Write headers
for col_idx, header in enumerate(headers, 1):
    cell = qs.cell(row=1, column=col_idx, value=header)
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Define Styles
zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
status_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
status_font = Font(name="Arial", size=10, bold=True, color="9C0006")

row_idx = 2
for idx, p in enumerate(presets, 1):
    vals = [
        idx, p['PresetName'], p['Strategy'], p['Timeframe'], p['PositionMode'],
        float(p['Risk_Pct']) if p['Risk_Pct'].replace('.', '', 1).isdigit() else p['Risk_Pct'],
        float(p['MinRR']) if p['MinRR'].replace('.', '', 1).isdigit() else p['MinRR'],
        int(p['MaxSpread']) if p['MaxSpread'].isdigit() else p['MaxSpread'],
        p['Session'], p['MomLookback'], p['MomThreshold'], p['Cooldown_Sec'],
        float(p['DailyLoss_Pct']) if p['DailyLoss_Pct'].replace('.', '', 1).isdigit() else p['DailyLoss_Pct'],
        float(p['MaxDD_Pct']) if p['MaxDD_Pct'].replace('.', '', 1).isdigit() else p['MaxDD_Pct'],
        p['TrailingStop'], p['UseBE'], "NOT TESTED"
    ]
    
    # Write descriptive values
    for col_idx, val in enumerate(vals, 1):
        cell = qs.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = Alignment(horizontal="center" if col_idx != 2 else "left", vertical="center")
        
        # Style Status
        if col_idx == 17:
            cell.fill = status_fill
            cell.font = status_font
        elif row_idx % 2 == 0 and col_idx < 17:
            # Apply zebra striping to descriptive columns
            cell.fill = zebra_fill

    # Write blank spaces for remaining columns
    for col_idx in range(18, len(headers) + 1):
        cell = qs.cell(row=row_idx, column=col_idx, value="")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx += 1

# Auto-fit columns
for col in qs.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    qs.column_dimensions[col_letter].width = max(max_len + 3, 10)

# Freeze headers
qs.freeze_panes = "A2"

wb.save(xlsx_path)
print(f"Successfully generated Test_Queue sheet in {xlsx_path}!")
